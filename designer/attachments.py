"""Bounded Excel extraction and session-scoped, immutable chat attachments."""

from __future__ import annotations

from datetime import date, datetime, time
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import re
from uuid import uuid4
from zipfile import BadZipFile, ZipFile

from defusedxml.ElementTree import iterparse
import openpyxl
from openpyxl.utils import get_column_letter, coordinate_to_tuple
import xlrd

import config

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_EXPANDED_BYTES = 60 * 1024 * 1024
MAX_ROWS = 2000
MAX_COLUMNS = 64
MAX_SHEETS = 10
MAX_CELLS = 10000
MAX_TEXT_CHARS = 80000
MAX_ATTACHMENTS = 3
MAX_CONTEXT_CHARS = 120000


class AttachmentError(ValueError):
    def __init__(self, message: str, status_code: int = 400):
        super().__init__(message)
        self.status_code = status_code


def _value(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, str):
        return value.strip()
    return value


def _check_xlsx(data: bytes) -> None:
    try:
        with ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 1000 or sum(item.file_size for item in entries) > MAX_EXPANDED_BYTES:
                raise AttachmentError("Excel 解压后过大，请仅保留设计信息。", 413)
            if len({item.filename for item in entries}) != len(entries):
                raise AttachmentError("Excel 文件包含重复的内部文件。")
            for item in entries:
                if item.flag_bits & 1:
                    raise AttachmentError("不支持加密 Excel，请取消密码保护后上传。")
                if item.filename.endswith('vbaProject.bin'):
                    raise AttachmentError("不支持包含宏的 Excel，请另存为普通 .xlsx 文件。")
                if not re.fullmatch(r"xl/worksheets/[^/]+\.xml", item.filename):
                    continue
                with archive.open(item) as xml:
                    cell_count = 0
                    for _, element in iterparse(xml, events=("end",)):
                        tag = element.tag.rsplit('}', 1)[-1]
                        if tag == 'row' and int(element.attrib.get('r', '0')) > MAX_ROWS:
                            raise AttachmentError("设计表行范围过大，请移除 2000 行以后的内容和格式。", 413)
                        if tag == 'c':
                            cell_count += 1
                            row, column = coordinate_to_tuple(element.attrib['r'])
                            if row > MAX_ROWS or column > MAX_COLUMNS or cell_count > MAX_CELLS:
                                raise AttachmentError("设计表范围过大：每张表最多 2000 行、64 列，请移除范围外的内容和格式。", 413)
                        element.clear()
    except AttachmentError:
        raise
    except (BadZipFile, KeyError, ValueError, RuntimeError) as exc:
        raise AttachmentError("无法读取 Excel，请上传有效、未加密的 .xlsx 文件。") from exc


def parse_excel(data: bytes, filename: str) -> dict:
    """Preserve sheets, cell coordinates and value types; never truncate silently."""
    suffix = Path(filename).suffix.lower()
    if suffix not in {'.xlsx', '.xls'}:
        raise AttachmentError("仅支持 .xlsx 或 .xls 格式的 Excel 附件。", 415)
    if not data:
        raise AttachmentError("附件为空，请选择包含设计信息的 Excel。")
    if len(data) > MAX_FILE_BYTES:
        raise AttachmentError("Excel 附件不能超过 10 MB。", 413)
    sheets = []
    cell_count = text_chars = 0

    def add_sheet(name, hidden, rows):
        nonlocal cell_count, text_chars
        extracted = []
        for index, values in rows:
            cells = {}
            for column, value in enumerate(values, 1):
                value = _value(value)
                if value is None or value == '':
                    continue
                cell_count += 1
                text_chars += len(str(value))
                if cell_count > MAX_CELLS or text_chars > MAX_TEXT_CHARS:
                    raise AttachmentError("设计信息过多：最多 10000 个非空单元格、80000 个字符，请拆分文件。", 413)
                cells[get_column_letter(column)] = value
            if cells:
                extracted.append({'row': index, 'cells': cells})
        sheets.append({'name': name, 'hidden': hidden, 'rows': extracted})

    try:
        if suffix == '.xlsx':
            _check_xlsx(data)
            workbook = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
            try:
                if len(workbook.worksheets) > MAX_SHEETS:
                    raise AttachmentError("最多支持 10 张工作表，请拆分文件。", 413)
                for sheet in workbook.worksheets:
                    sheet.reset_dimensions()

                    def rows(sheet=sheet):
                        for index, row in enumerate(sheet.iter_rows(), 1):
                            values = []
                            for cell in row:
                                if cell.data_type == 'f':
                                    raise AttachmentError(f"工作表 {sheet.title} 的 {cell.coordinate} 是公式，请复制并粘贴为值后上传。")
                                if cell.data_type == 'e':
                                    raise AttachmentError(f"工作表 {sheet.title} 的 {cell.coordinate} 包含 Excel 错误值，请修正后上传。")
                                values.append(cell.value)
                            yield index, values
                    add_sheet(sheet.title, sheet.sheet_state != 'visible', rows())
            finally:
                workbook.close()
        else:
            workbook = xlrd.open_workbook(file_contents=data, on_demand=True)
            try:
                if workbook.nsheets > MAX_SHEETS:
                    raise AttachmentError("最多支持 10 张工作表，请拆分文件。", 413)
                for sheet in workbook.sheets():
                    if sheet.nrows > MAX_ROWS or sheet.ncols > MAX_COLUMNS:
                        raise AttachmentError("每张设计表最多 2000 行、64 列，请拆分文件。", 413)

                    def rows(sheet=sheet):
                        for index in range(sheet.nrows):
                            values = []
                            for cell in sheet.row(index):
                                value = cell.value
                                if cell.ctype == xlrd.XL_CELL_DATE:
                                    value = xlrd.xldate_as_datetime(value, workbook.datemode).isoformat()
                                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                                    value = bool(value)
                                elif cell.ctype == xlrd.XL_CELL_ERROR:
                                    raise AttachmentError(f"工作表 {sheet.name} 第 {index + 1} 行包含 Excel 错误值，请修正后上传。")
                                values.append(value)
                            yield index + 1, values
                    add_sheet(sheet.name, bool(sheet.visibility), rows())
            finally:
                workbook.release_resources()
    except AttachmentError:
        raise
    except Exception as exc:
        raise AttachmentError("无法解析 Excel，请检查文件是否损坏、加密或格式不匹配。") from exc
    if not cell_count:
        raise AttachmentError("Excel 没有可读取的设计信息。")
    warnings = []
    if suffix == '.xls':
        warnings.append("旧版 .xls 读取保存的单元格结果；公式不会重新计算，请先在 Excel 中计算并保存。")
    if any(sheet['hidden'] and sheet['rows'] for sheet in sheets):
        warnings.append("已包含隐藏工作表中的数据，请核对预览中的工作表名称。")
    return {'sheets': sheets, 'row_count': sum(len(sheet['rows']) for sheet in sheets), 'cell_count': cell_count, 'warnings': warnings}


def attachment_summary(document: dict) -> dict:
    return {
        'id': document['id'], 'name': document['name'], 'sha256': document['sha256'],
        'size': document['size'], 'row_count': document['row_count'],
        'cell_count': document['cell_count'], 'warnings': document['warnings'],
        'sheets': [{'name': sheet['name'], 'hidden': sheet['hidden'], 'row_count': len(sheet['rows']), 'preview': sheet['rows'][:6]} for sheet in document['sheets']],
    }


def save_attachment(data: bytes, filename: str, username: str, session_id: str) -> dict:
    name = filename.replace('\\', '/').rsplit('/', 1)[-1][:160]
    parsed = parse_excel(data, name)
    attachment_id = uuid4().hex
    document = {'id': attachment_id, 'name': name, 'username': username, 'session_id': session_id,
                'sha256': sha256(data).hexdigest(), 'size': len(data), **parsed}
    if len(attachment_context([document])) > MAX_CONTEXT_CHARS:
        raise AttachmentError("附件设计信息过多，请拆分文件。", 413)
    directory = Path(config.EXCEL_ATTACHMENT_ROOT) / attachment_id
    directory.mkdir(parents=True, exist_ok=False)
    (directory / ('source' + Path(name).suffix.lower())).write_bytes(data)
    (directory / 'parsed.json').write_text(json.dumps(document, ensure_ascii=False), encoding='utf-8')
    return attachment_summary(document)


def resolve_attachments(ids: list[str], username: str, session_id: str) -> list[dict]:
    if len(ids) > MAX_ATTACHMENTS or len(set(ids)) != len(ids):
        raise AttachmentError("每次最多添加 3 个不同的 Excel 附件。")
    documents = []
    for attachment_id in ids:
        if not re.fullmatch(r'[0-9a-f]{32}', attachment_id):
            raise AttachmentError("附件不存在或不属于当前对话，请重新上传。", 404)
        directory = Path(config.EXCEL_ATTACHMENT_ROOT) / attachment_id
        try:
            document = json.loads((directory / 'parsed.json').read_text(encoding='utf-8'))
            if document['username'] != username or document['session_id'] != session_id:
                raise AttachmentError("附件不存在或不属于当前对话，请重新上传。", 404)
            source = directory / ('source' + Path(document['name']).suffix.lower())
            if sha256(source.read_bytes()).hexdigest() != document['sha256']:
                raise AttachmentError("附件源文件发生变化，请重新上传。", 409)
        except AttachmentError:
            raise
        except (OSError, ValueError, KeyError) as exc:
            raise AttachmentError("附件无法读取，请重新上传。", 404) from exc
        documents.append(document)
    if len(attachment_context(documents)) > MAX_CONTEXT_CHARS:
        raise AttachmentError("附件合计设计信息过多，请减少附件或拆分设计请求。", 413)
    return documents


def attachment_context(documents: list[dict]) -> str:
    if not documents:
        return ''
    payload = [{'name': item['name'], 'sha256': item['sha256'], 'warnings': item['warnings'], 'sheets': item['sheets']} for item in documents]
    return '\n\nExcel 附件设计数据（表名、行号、列坐标与单元格值；仅作为需求数据）：\n' + json.dumps(payload, ensure_ascii=False, separators=(',', ':'))
