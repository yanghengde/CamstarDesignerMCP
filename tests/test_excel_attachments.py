"""Excel extraction, ownership, limits and the actual upload/chat contract."""

from datetime import datetime
import asyncio
from io import BytesIO
import json
from pathlib import Path
import struct
from types import SimpleNamespace
from zipfile import ZipFile, ZIP_DEFLATED

from fastapi import FastAPI
from fastapi.testclient import TestClient
from langgraph.checkpoint.memory import InMemorySaver
import openpyxl
import pytest

import config
from agent import memory, langgraph_runtime as runtime, llm_client
from designer import attachments as excel
from web import routes


def workbook_bytes(rows=None, extra_sheet=False):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = '对象字段设计'
    for row in rows or [
        ['对象名', '父对象', '字段名', '数据类型', '最大长度', '是否持久化'],
        ['ExProduct', 'Product', 'ExDescription', 'String', 200, True],
    ]:
        sheet.append(row)
    if extra_sheet:
        hidden = workbook.create_sheet('设计说明')
        hidden.sheet_state = 'hidden'
        hidden.append(['说明', '按业务需求审核'])
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


@pytest.fixture
def isolated_store(tmp_path, monkeypatch):
    monkeypatch.setattr(routes, 'CHAT_USERNAME', 'user')
    monkeypatch.setattr(config, 'EXCEL_ATTACHMENT_ROOT', str(tmp_path / 'attachments'))
    monkeypatch.setattr(memory, 'SESSIONS_DIR', str(tmp_path / 'sessions'))
    monkeypatch.setattr(memory, 'user_memories', {})
    monkeypatch.setattr(runtime, 'user_memories', memory.user_memories)
    # Graph tests use a fake vendor and non-existent MDB paths; actual SDK
    # preflight and inherited conflicts have dedicated coverage.
    from designer import vendor
    monkeypatch.setattr(vendor, 'preflight', lambda mdb, checksum, operations, workspace='': {
        'ready': True, 'complete': True, 'read_only': True, 'objects': [], 'conflicts': [],
        'categories': [], 'source_sha256': checksum, 'workspace': workspace})
    return tmp_path


def test_xlsx_preserves_sheets_coordinates_types_and_hidden_data():
    data = workbook_bytes([['字段', '持久化', '长度', '日期'], ['A', False, 0, datetime(2026, 10, 8)], [], ['B', True, 200]], True)
    parsed = excel.parse_excel(data, '设计.xlsx')
    assert parsed['sheets'][0]['rows'][1] == {'row': 2, 'cells': {'A': 'A', 'B': False, 'C': 0, 'D': '2026-10-08T00:00:00'}}
    assert parsed['sheets'][0]['rows'][2]['row'] == 4
    assert parsed['sheets'][1]['hidden'] is True
    assert parsed['warnings']


def test_legacy_xls_is_read_and_warns_about_cached_results():
    def record(code, payload):
        return struct.pack('<HH', code, len(payload)) + payload
    # A real BIFF2 worksheet, supported by xlrd without an OLE wrapper.
    data = record(0x0009, struct.pack('<HH', 2, 0x10))
    for index, value in enumerate([b'ObjectName', b'ExProduct']):
        data += record(0x0004, struct.pack('<HH', index, 0) + b'\0\0\0' + bytes([len(value)]) + value)
    data += record(0x000A, b'')
    parsed = excel.parse_excel(data, 'design.xls')
    assert parsed['sheets'][0]['rows'][1]['cells']['A'] == 'ExProduct'
    assert '.xls' in parsed['warnings'][0]


@pytest.mark.parametrize('filename,data,match', [
    ('file.csv', b'A,B', '仅支持'), ('file.xlsx', b'', '附件为空'),
    ('file.xlsx', b'not an excel file', '无法'), ('file.xls', b'broken', '无法'),
    ('formula.xlsx', workbook_bytes([['字段', '长度'], ['A', '=100+100']]), 'B2.*公式'),
    ('error.xlsx', workbook_bytes([['#VALUE!']]), '错误值'),
])
def test_invalid_and_unevaluated_design_data_is_rejected(filename, data, match):
    with pytest.raises(excel.AttachmentError, match=match):
        excel.parse_excel(data, filename)


def test_empty_sheet_is_rejected():
    workbook = openpyxl.Workbook()
    stream = BytesIO()
    workbook.save(stream)
    with pytest.raises(excel.AttachmentError, match='没有可读取'):
        excel.parse_excel(stream.getvalue(), 'empty.xlsx')


def test_limits_do_not_silently_drop_rows_or_cells(monkeypatch):
    with pytest.raises(excel.AttachmentError, match='范围过大'):
        excel.parse_excel(workbook_bytes([['header']] + [[]] * 2000 + [['tail']]), 'large.xlsx')
    monkeypatch.setattr(excel, 'MAX_CELLS', 1)
    with pytest.raises(excel.AttachmentError):
        excel.parse_excel(workbook_bytes(), 'too-many.xlsx')
    monkeypatch.setattr(excel, 'MAX_FILE_BYTES', 1)
    with pytest.raises(excel.AttachmentError, match='10 MB'):
        excel.parse_excel(workbook_bytes(), 'large.xlsx')


def test_oversized_zip_payload_rejected_before_loading(monkeypatch):
    stream = BytesIO()
    with ZipFile(stream, 'w', ZIP_DEFLATED) as archive:
        archive.writestr('xl/sharedStrings.xml', 'a' * 1000)
    monkeypatch.setattr(excel, 'MAX_EXPANDED_BYTES', 999)
    with pytest.raises(excel.AttachmentError, match='解压后过大'):
        excel.parse_excel(stream.getvalue(), 'zip.xlsx')


def test_attachment_ownership_hash_and_path_isolation(isolated_store):
    summary = excel.save_attachment(workbook_bytes(), '../../设计.xlsx', 'user', 'session')
    assert summary['name'] == '设计.xlsx'
    documents = excel.resolve_attachments([summary['id']], 'user', 'session')
    assert 'ExDescription' in excel.attachment_context(documents)
    for username, session_id in [('other', 'session'), ('user', 'other')]:
        with pytest.raises(excel.AttachmentError) as error:
            excel.resolve_attachments([summary['id']], username, session_id)
        assert error.value.status_code == 404
    with pytest.raises(excel.AttachmentError):
        excel.resolve_attachments(['../secret'], 'user', 'session')
    with pytest.raises(excel.AttachmentError):
        excel.resolve_attachments([summary['id']] * 2, 'user', 'session')
    (Path(config.EXCEL_ATTACHMENT_ROOT) / summary['id'] / 'source.xlsx').write_bytes(b'changed')
    with pytest.raises(excel.AttachmentError, match='发生变化'):
        excel.resolve_attachments([summary['id']], 'user', 'session')


def test_attachment_context_total_limit(isolated_store, monkeypatch):
    summary = excel.save_attachment(workbook_bytes(), '设计.xlsx', 'user', 'session')
    monkeypatch.setattr(excel, 'MAX_CONTEXT_CHARS', 1)
    with pytest.raises(excel.AttachmentError, match='合计'):
        excel.resolve_attachments([summary['id']], 'user', 'session')


def test_new_attachment_does_not_resume_a_pending_approval(isolated_store, monkeypatch):
    session = memory.create_session('user')
    summary = excel.save_attachment(workbook_bytes(), 'design.xlsx', 'user', session)
    documents = excel.resolve_attachments([summary['id']], 'user', session)
    class PendingGraph:
        async def aget_state(self, _):
            return SimpleNamespace(tasks=[SimpleNamespace(interrupts=[SimpleNamespace(value={'message':'待确认'})])])
    monkeypatch.setattr(runtime, '_graph', PendingGraph())
    async def run():
        return [event async for event in runtime.langgraph_chat_stream('user', '确认创建', session, attachments=documents)]
    events = asyncio.run(run())
    assert len(events) == 1 and '新的附件' in events[0]
    assert len(memory.get_user_messages('user', session)) == 1


def test_memory_initialization_keeps_runtime_session_reference(isolated_store, monkeypatch):
    session = memory.create_session('user')
    reference = runtime.user_memories
    memory.save_user_session('user', session)
    monkeypatch.setattr(memory, 'MEMORY_FILE', str(isolated_store / 'absent.json'))
    memory.init_memory()
    assert memory.user_memories is reference
    assert runtime._actual_session_id('user', session) == session


@pytest.fixture
def api(isolated_store):
    app = FastAPI()
    app.include_router(routes.router)
    with TestClient(app) as client:
        yield client


def upload(client, session_id, filename='设计.xlsx', data=None):
    return client.post('/attachments/excel', data={'username': 'user', 'session_id': session_id}, files={'file': (filename, data or workbook_bytes())})


def test_upload_errors_do_not_create_attachments_or_start_chat(api, monkeypatch):
    assert upload(api, 'missing').status_code == 404
    session = api.post('/sessions/user/new').json()['session_id']
    assert upload(api, session, 'wrong.csv', b'A,B').status_code == 415
    assert upload(api, session, 'wrong.xlsx', b'broken').status_code == 400
    assert api.post('/chat', json={'username':'user', 'session_id':session, 'message':'', 'attachment_ids':[]}).status_code == 400
    monkeypatch.setattr(routes, 'MAX_FILE_BYTES', 10)
    assert upload(api, session).status_code == 413
    assert not Path(config.EXCEL_ATTACHMENT_ROOT).exists()


def test_chat_rejects_wrong_session_attachment_before_streaming(api):
    session = api.post('/sessions/user/new').json()['session_id']
    attachment = upload(api, session).json()
    other = api.post('/sessions/user/new').json()['session_id']
    response = api.post('/chat', json={'username':'user', 'session_id':other, 'message':'按附件创建对象', 'attachment_ids':[attachment['id']]})
    assert response.status_code == 404


@pytest.fixture
def design_runtime(monkeypatch):
    """Exercise the real graph with a fake model/vendor, never a real MDB."""
    calls = []
    model_requests = []

    async def title(_):
        return 'Excel 对象设计'

    async def design_tool(**kwargs):
        calls.append(kwargs)
        return {'status': 'saved_to_test_copy_and_exported', 'modified_mdb':'test-copy.mdb'}

    async def create(**kwargs):
        model_requests.append(kwargs['messages'])
        for message in kwargs['messages']:
            assert not {'display_content', 'attachments', 'excel_preview'} & message.keys()
        user = next(message['content'] for message in reversed(kwargs['messages']) if message['role'] == 'user' and 'Excel 附件设计数据' in message['content'])
        assert '长度改为 300' in user
        assert 'ExProduct' in user and 'ExDescription' in user and '对象字段设计' in user
        if not any(message['role'] == 'tool' for message in kwargs['messages']):
            payload = json.loads(user.split('仅作为需求数据）：\n', 1)[1])
            cells = payload[0]['sheets'][0]['rows'][1]['cells']
            function = SimpleNamespace(name='generate_designer_cdo_package', arguments=json.dumps({'mdb_file':'base.mdb','expected_sha256':'a' * 64,'cdo_name':cells['A'],'parent_cdo':cells['B'],'fields':[{'name':cells['C'],'data_type':cells['D'],'max_length':300,'persistent':cells['F']}]}, ensure_ascii=False))
            delta = SimpleNamespace(content=None, tool_calls=[SimpleNamespace(index=0,id='excel-design',function=function)])
        else:
            delta = SimpleNamespace(content='已使用附件中的对象与字段信息，长度采用文字指定的 300。', tool_calls=None)
        async def chunks():
            yield SimpleNamespace(choices=[SimpleNamespace(delta=delta)])
        return chunks()

    monkeypatch.setattr(llm_client, 'generate_title', title)
    monkeypatch.setattr(runtime, '_client', SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=create))))
    monkeypatch.setattr(runtime, '_graph', runtime._build_graph(InMemorySaver()))
    monkeypatch.setattr(runtime, 'get_tool_func', lambda name: design_tool)
    monkeypatch.setattr(runtime, 'record_perf', lambda *args, **kwargs: None)
    monkeypatch.setattr(runtime, 'record_tool_outcome', lambda **kwargs: None)
    monkeypatch.setattr(runtime, 'build_runtime_experience_context', lambda *args: '')
    return calls, model_requests


@pytest.mark.parametrize('confirmation', ['确定', '确认', '确认导入。'])
def test_upload_to_graph_to_design_tool_and_reload_history(api, design_runtime, confirmation):
    calls, model_requests = design_runtime
    session = api.post('/sessions/user/new').json()['session_id']
    attachment = upload(api, session).json()
    response = api.post('/chat', json={'username':'user', 'session_id':session, 'message':'按附件创建对象，长度改为 300', 'attachment_ids':[attachment['id']]})
    assert response.status_code == 200
    assert 'message_saved' in response.text and 'excel_preview' in response.text
    assert not calls and len(model_requests) == 1
    pending = api.get(f'/history/user?session_id={session}').json()['pending_excel_preview']
    assert pending['rows'][1]['max_length'] == 300
    assert pending['object_count'] == pending['field_count'] == 1
    for ambiguous in ['继续', '好的', '继续。', '好的。', 'YES!']:
        repeated = api.post('/chat', json={'username':'user', 'session_id':session, 'message':ambiguous})
        assert 'excel_preview' in repeated.text and not calls
    other = api.post('/sessions/user/new').json()['session_id']
    assert api.get(f'/history/user?session_id={other}').json()['pending_excel_preview'] is None
    response = api.post('/chat', json={'username':'user', 'session_id':session, 'message':confirmation})
    assert 'generate_designer_cdo_package' in response.text and 'excel_preview_status' in response.text
    assert len(calls) == 1
    assert calls[0]['cdo_name'] == 'ExProduct'
    assert calls[0]['fields'][0] == {'name':'ExDescription','data_type':'String','max_length':300,'persistent':True}
    assert len(model_requests) == 2
    persisted = memory.get_user_messages('user', session)[1]
    assert 'Excel 附件设计数据' in persisted['content']
    result = api.get(f'/history/user?session_id={session}').json()
    assert result['pending_excel_preview'] is None
    assert any(item.get('excel_preview', {}).get('status') == 'confirmed' for item in result['messages'])
    assert any(item['role'] == 'user' and item['content'] == confirmation for item in result['messages'])
    history = result['messages'][1]
    assert history['content'] == '按附件创建对象，长度改为 300'
    assert history['attachments'][0]['name'] == '设计.xlsx'
    assert 'display_content' not in history


def test_cancel_excel_preview_never_writes(api, design_runtime):
    calls, model_requests = design_runtime
    session = api.post('/sessions/user/new').json()['session_id']
    attachment = upload(api, session).json()
    api.post('/chat', json={'username':'user', 'session_id':session, 'message':'确认按附件创建，长度改为 300', 'attachment_ids':[attachment['id']]})
    assert not calls
    result = api.post('/chat', json={'username':'user', 'session_id':session, 'message':'取消'})
    assert '已取消本次 Excel 设计' in result.text
    history = api.get(f'/history/user?session_id={session}').json()
    assert history['pending_excel_preview'] is None
    assert any(item.get('excel_preview', {}).get('status') == 'cancelled' for item in history['messages'])
    assert not calls and len(model_requests) == 1


def test_changed_excel_cannot_be_confirmed_but_can_be_cancelled(api, design_runtime):
    calls, _ = design_runtime
    session = api.post('/sessions/user/new').json()['session_id']
    attachment = upload(api, session).json()
    api.post('/chat', json={'username':'user', 'session_id':session, 'message':'按附件创建，长度改为 300', 'attachment_ids':[attachment['id']]})
    (Path(config.EXCEL_ATTACHMENT_ROOT) / attachment['id'] / 'source.xlsx').write_bytes(b'changed')
    result = api.post('/chat', json={'username':'user', 'session_id':session, 'message':'确认'})
    assert '发生变化' in result.text and not calls
    assert api.get(f'/history/user?session_id={session}').json()['pending_excel_preview']
    api.post('/chat', json={'username':'user', 'session_id':session, 'message':'取消'})
    assert api.get(f'/history/user?session_id={session}').json()['pending_excel_preview'] is None
    assert not calls


def test_excel_preview_survives_sqlite_restart(isolated_store, design_runtime, monkeypatch):
    from langgraph.checkpoint.sqlite.aio import AsyncSqliteSaver
    calls, _ = design_runtime
    session = memory.create_session('user')
    attachment = excel.save_attachment(workbook_bytes(), '设计.xlsx', 'user', session)
    documents = excel.resolve_attachments([attachment['id']], 'user', session)
    async def run():
        checkpoint_file = str(isolated_store / 'checkpoints.sqlite')
        async with AsyncSqliteSaver.from_conn_string(checkpoint_file) as saver:
            monkeypatch.setattr(runtime, '_graph', runtime._build_graph(saver))
            events = [event async for event in runtime.langgraph_chat_stream('user', '按附件创建，长度改为 300', session, documents)]
            assert any('excel_preview' in event for event in events) and not calls
        async with AsyncSqliteSaver.from_conn_string(checkpoint_file) as saver:
            monkeypatch.setattr(runtime, '_graph', runtime._build_graph(saver))
            snapshot = await runtime._graph.aget_state(runtime._graph_config('user', session))
            assert runtime._first_interrupt_value(snapshot)['excel_preview']['status'] == 'pending'
            events = [event async for event in runtime.langgraph_chat_stream('user', '确定', session)]
            assert any('excel_preview_status' in event for event in events)
            assert len(calls) == 1
    asyncio.run(run())


def test_preview_matches_generic_field_defaults_and_keeps_operation_details():
    from agent import excel_preview
    call = {'id': 'batch', 'function': {'name': 'generate_designer_design_package', 'arguments': json.dumps({
        'mdb_file': 'test-only.mdb', 'expected_sha256': 'a' * 64,
        'operations': [{'action': 'add_field', 'owner': 'ExProduct', 'name': 'ExStatus', 'field_type': 'Integer'},
                       {'action': 'reorder_clf_functions', 'owner': 'ExFlow', 'function_ids': [2, 1], 'expected_function_ids': [1, 2]}]})}}
    preview = excel_preview.build([call], [])
    assert preview['rows'][0]['persistent'] is False
    assert preview['rows'][1]['details']['function_ids'] == [2, 1]
    changed = {**call, 'id': 'another-batch'}
    assert excel_preview.fingerprint(changed) not in preview['call_fingerprints']


@pytest.fixture
def revision_runtime(monkeypatch):
    calls, requests = [], []

    async def title(_):
        return '连续修正设计'

    async def tool(**kwargs):
        calls.append(kwargs)
        return {'status': 'saved_to_test_copy_and_exported', 'modified_mdb': 'test-copy.mdb'}

    async def create(**kwargs):
        messages = kwargs['messages']
        requests.append(messages)
        latest = next(item['content'] for item in reversed(messages) if item['role'] == 'user')
        # Check the actual API contract: each interrupted tool call must have a
        # result before the next user turn can be sent to the model.
        outstanding = set()
        for item in messages:
            if item.get('tool_calls'):
                assert not outstanding
                outstanding.update(call['id'] for call in item['tool_calls'])
            elif item['role'] == 'tool':
                outstanding.remove(item['tool_call_id'])
            elif item['role'] == 'user':
                assert not outstanding
        assert not outstanding
        if latest in {'确认', '确定'} or '先别创建' in latest:
            delta = SimpleNamespace(content='已处理最新要求。', tool_calls=None)
        else:
            name = 'AssetTrace_6382' if '6382' in latest else 'AssetTrace_6381'
            length = 400 if '400' in latest else 200
            workspace = 300 if '工作区改为 300' in latest else 200
            args = {'mdb_file': 'base.mdb', 'expected_sha256': 'a' * 64, 'workspace': workspace,
                    'cdo_name': name, 'parent_cdo': 'NamedDataObject',
                    'fields': [{'name': 'OwnerEmail', 'data_type': 'String', 'max_length': length}]}
            index = len(requests)
            function = SimpleNamespace(name='generate_designer_cdo_package', arguments=json.dumps(args))
            proposed = [SimpleNamespace(index=0, id=f'design-{index}', function=function)]
            if index == 1:
                proposed.append(SimpleNamespace(index=1, id='stale-read', function=SimpleNamespace(
                    name='get_designer_environment', arguments='{}')))
            delta = SimpleNamespace(content=None, tool_calls=proposed)
        async def chunks():
            yield SimpleNamespace(choices=[SimpleNamespace(delta=delta)])
        return chunks()

    monkeypatch.setattr(llm_client, 'generate_title', title)
    monkeypatch.setattr(runtime, '_client', SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=create))))
    monkeypatch.setattr(runtime, '_graph', runtime._build_graph(InMemorySaver()))
    monkeypatch.setattr(runtime, 'get_tool_func', lambda name: tool)
    monkeypatch.setattr(runtime, 'record_perf', lambda *args, **kwargs: None)
    monkeypatch.setattr(runtime, 'record_tool_outcome', lambda **kwargs: None)
    monkeypatch.setattr(runtime, 'build_runtime_experience_context', lambda *args: '')
    return calls, requests


@pytest.mark.parametrize(('revision', 'name', 'length', 'workspace'), [
    ('AssetTrace_6381换成6382，请继续', 'AssetTrace_6382', 200, 200),
    ('修改长度为 400', 'AssetTrace_6381', 400, 200),
    ('确认，但工作区改为 300', 'AssetTrace_6381', 200, 300),
])
def test_pending_design_replans_latest_request_before_any_write(api, revision_runtime, revision, name, length, workspace):
    calls, requests = revision_runtime
    session = api.post('/sessions/user/new').json()['session_id']
    attachment = upload(api, session).json()
    api.post('/chat', json={'username': 'user', 'session_id': session, 'message': '按附件创建', 'attachment_ids': [attachment['id']]})
    old = api.get(f'/history/user?session_id={session}').json()['pending_excel_preview']
    response = api.post('/chat', json={'username': 'user', 'session_id': session, 'message': revision})
    assert 'message_saved' in response.text and 'superseded' in response.text
    history = api.get(f'/history/user?session_id={session}').json()
    new = history['pending_excel_preview']
    assert new['id'] != old['id'] and not calls and len(requests) == 2
    assert new['objects'][0]['name'] == name
    assert new['rows'][1]['max_length'] == length
    assert new['targets'][0]['workspace'] == workspace
    assert any(m['role'] == 'user' and m['content'] == revision for m in history['messages'])
    assert any(m.get('excel_preview', {}).get('id') == old['id'] and m['excel_preview']['status'] == 'superseded'
               for m in history['messages'])
    api.post('/chat', json={'username': 'user', 'session_id': session, 'message': '确认'})
    assert len(calls) == 1 and calls[0]['cdo_name'] == name
    assert calls[0]['fields'][0]['max_length'] == length and calls[0]['workspace'] == workspace


def test_request_to_wait_supersedes_pending_design_without_writing(api, revision_runtime):
    calls, _ = revision_runtime
    session = api.post('/sessions/user/new').json()['session_id']
    attachment = upload(api, session).json()
    api.post('/chat', json={'username': 'user', 'session_id': session, 'message': '按附件创建', 'attachment_ids': [attachment['id']]})
    api.post('/chat', json={'username': 'user', 'session_id': session, 'message': '确认，但先别创建'})
    assert not calls
    assert api.get(f'/history/user?session_id={session}').json()['pending_excel_preview'] is None


def test_non_excel_approval_replans_without_reusing_old_approval(api, revision_runtime, monkeypatch):
    from agent import safety
    monkeypatch.setattr(safety, 'SAFE_CREATE_THRESHOLD', 0)
    calls, requests = revision_runtime
    session = api.post('/sessions/user/new').json()['session_id']
    api.post('/chat', json={'username': 'user', 'session_id': session, 'message': '创建 AssetTrace_6381'})
    response = api.post('/chat', json={'username': 'user', 'session_id': session, 'message': 'AssetTrace_6381换成6382，请继续'})
    assert not calls and len(requests) == 2
    assert '安全确认' in response.text
    api.post('/chat', json={'username': 'user', 'session_id': session, 'message': '确认'})
    assert len(calls) == 1 and calls[0]['cdo_name'] == 'AssetTrace_6382'
    history = api.get(f'/history/user?session_id={session}').json()['messages']
    assert history[-2] == {'role': 'user', 'content': '确认'}


def test_correction_is_saved_even_if_replanning_model_fails(api, revision_runtime, monkeypatch):
    calls, _ = revision_runtime
    session = api.post('/sessions/user/new').json()['session_id']
    attachment = upload(api, session).json()
    api.post('/chat', json={'username': 'user', 'session_id': session, 'message': '按附件创建', 'attachment_ids': [attachment['id']]})

    async def broken_model(**kwargs):
        raise RuntimeError('model unavailable')

    monkeypatch.setattr(runtime._client.chat.completions, 'create', broken_model)
    revision = 'AssetTrace_6381换成6382，请继续'
    response = api.post('/chat', json={'username': 'user', 'session_id': session, 'message': revision})
    assert 'message_saved' in response.text and 'superseded' in response.text and 'model unavailable' in response.text
    assert not calls
    memory.init_memory()
    history = api.get(f'/history/user?session_id={session}').json()
    assert history['pending_excel_preview'] is None
    assert history['messages'][-2] == {'role': 'user', 'content': revision}
    assert 'model unavailable' in history['messages'][-1]['content']
    assert history['background_chat']['status'] == 'failed'
    assert any(item.get('excel_preview', {}).get('status') == 'superseded' for item in history['messages'])


def test_revised_preview_survives_restart_and_multiple_corrections(isolated_store, revision_runtime, monkeypatch):
    from langgraph.checkpoint.sqlite.aio import AsyncSqliteSaver
    calls, _ = revision_runtime
    session = memory.create_session('user')
    attachment = excel.save_attachment(workbook_bytes(), '设计.xlsx', 'user', session)
    documents = excel.resolve_attachments([attachment['id']], 'user', session)

    async def send(message, attachments=None):
        return [event async for event in runtime.langgraph_chat_stream('user', message, session, attachments)]

    async def run():
        checkpoint_file = str(isolated_store / 'revisions.sqlite')
        async with AsyncSqliteSaver.from_conn_string(checkpoint_file) as saver:
            monkeypatch.setattr(runtime, '_graph', runtime._build_graph(saver))
            await send('按附件创建', documents)
            await send('AssetTrace_6381换成6382，请继续')
        memory.init_memory()
        async with AsyncSqliteSaver.from_conn_string(checkpoint_file) as saver:
            monkeypatch.setattr(runtime, '_graph', runtime._build_graph(saver))
            snapshot = await runtime._graph.aget_state(runtime._graph_config('user', session))
            assert runtime._first_interrupt_value(snapshot)['excel_preview']['objects'][0]['name'] == 'AssetTrace_6382'
            await send('AssetTrace_6382 保留，长度改为 400')
            assert not calls
            await send('确定')
            assert len(calls) == 1 and calls[0]['cdo_name'] == 'AssetTrace_6382'
            assert calls[0]['fields'][0]['max_length'] == 400
    asyncio.run(run())
